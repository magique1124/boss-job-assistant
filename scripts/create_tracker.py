#!/usr/bin/env python3
"""创建投递追踪表（xlsx）。已存在则不改动，直接退出。

用法: python3 create_tracker.py <输出路径.xlsx>
"""
import os
import sys

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

COLUMNS = [
    ("序号", 6),
    ("公司", 20),
    ("岗位", 22),
    ("薪资", 14),
    ("JD 链接", 40),
    ("是否投递", 10),
    ("是否回复", 10),
    ("是否约面", 10),
    ("面试情况", 24),
    ("备注", 30),
    ("更新时间", 12),
]

YES_NO = '"是,否"'
APPLY_STATUS = '"已投递,待投递,未投递"'
INTERVIEW_STATUS = '"未开始,待面试,通过,未通过,待定"'


def create(path: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "投递追踪"

    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill("solid", fgColor="4472C4")
    center = Alignment(horizontal="center", vertical="center")

    for idx, (name, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=idx, value=name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center
        ws.column_dimensions[get_column_letter(idx)].width = width

    ws.freeze_panes = "A2"

    validations = [
        ("F", APPLY_STATUS),
        ("G", YES_NO),
        ("H", YES_NO),
        ("I", INTERVIEW_STATUS),
    ]
    for col, formula in validations:
        dv = DataValidation(type="list", formula1=formula, allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{col}2:{col}5000")

    wb.save(path)


def main() -> int:
    if len(sys.argv) != 2:
        print(__doc__)
        return 1
    path = sys.argv[1]
    if os.path.exists(path):
        try:
            load_workbook(path)
        except Exception as exc:  # noqa: BLE001
            print(f"ERROR: 文件已存在但不是有效的 xlsx: {exc}")
            return 1
        print(f"EXISTS: {path}")
        return 0
    create(path)
    print(f"CREATED: {path}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
